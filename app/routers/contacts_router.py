import io
from fastapi import APIRouter, Request, Form, UploadFile, File, Depends
from fastapi.responses import RedirectResponse, HTMLResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from ..db import get_db
from ..models import Contact
from ..auth import require_login

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")


def normalize_phone(raw):
    digits = "".join(ch for ch in str(raw) if ch.isdigit())
    if len(digits) == 10:
        digits = "91" + digits
    return digits


@router.get("/contacts", response_class=HTMLResponse)
def list_contacts(request: Request, db: Session = Depends(get_db)):
    if (resp := require_login(request)):
        return resp
    contacts = db.query(Contact).order_by(Contact.created_at.desc()).all()
    return templates.TemplateResponse(
        "contacts.html",
        {"request": request, "logged_in": True, "contacts": contacts},
    )


@router.post("/contacts/add")
def add_contact(request: Request, phone: str = Form(...), name: str = Form(""), db: Session = Depends(get_db)):
    if (resp := require_login(request)):
        return resp
    phone_n = normalize_phone(phone)
    existing = db.query(Contact).filter(Contact.phone == phone_n).first()
    if existing:
        existing.name = name
    else:
        db.add(Contact(phone=phone_n, name=name))
    db.commit()
    return RedirectResponse(url="/contacts", status_code=302)


@router.post("/contacts/delete/{contact_id}")
def delete_contact(request: Request, contact_id: int, db: Session = Depends(get_db)):
    if (resp := require_login(request)):
        return resp
    c = db.query(Contact).filter(Contact.id == contact_id).first()
    if c:
        db.delete(c)
        db.commit()
    return RedirectResponse(url="/contacts", status_code=302)


@router.post("/contacts/import")
async def import_contacts(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    if (resp := require_login(request)):
        return resp
    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]

    def col(name):
        return headers.index(name) + 1 if name in headers else None

    phone_col = col("phone")
    name_col = col("name")
    added, updated = 0, 0
    if phone_col:
        for r in range(2, ws.max_row + 1):
            phone_val = ws.cell(row=r, column=phone_col).value
            if not phone_val:
                continue
            phone_n = normalize_phone(phone_val)
            name_val = ws.cell(row=r, column=name_col).value if name_col else ""
            existing = db.query(Contact).filter(Contact.phone == phone_n).first()
            if existing:
                existing.name = str(name_val).strip() if name_val else existing.name
                updated += 1
            else:
                db.add(Contact(phone=phone_n, name=str(name_val).strip() if name_val else ""))
                added += 1
        db.commit()
    return RedirectResponse(url="/contacts", status_code=302)
